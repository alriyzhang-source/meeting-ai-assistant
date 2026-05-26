"""AI 回答模块 — 检测问题并调用 Claude 生成回答"""

import os
import queue
import threading
from anthropic import Anthropic
from config import CLAUDE_API_KEY, CLAUDE_MODEL, MATERIALS_DIR, QUESTION_KEYWORDS, QUESTION_PUNCTUATION


def _read_txt(fpath: str) -> str:
    with open(fpath, "r", encoding="utf-8") as f:
        return f.read()


def _read_pdf(fpath: str) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(fpath)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages)


def _read_pptx(fpath: str) -> str:
    from pptx import Presentation
    prs = Presentation(fpath)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            slides.append(f"[幻灯片 {i}] " + " | ".join(texts))
    return "\n".join(slides)


def _read_docx(fpath: str) -> str:
    from docx import Document
    doc = Document(fpath)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


READERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".pdf": _read_pdf,
    ".pptx": _read_pptx,
    ".docx": _read_docx,
}


def load_materials() -> str:
    """加载 materials/ 目录下的所有支持格式文件"""
    materials = []
    if not os.path.isdir(MATERIALS_DIR):
        print(f"[AI] 材料目录不存在: {MATERIALS_DIR}")
        return ""

    for fname in sorted(os.listdir(MATERIALS_DIR)):
        fpath = os.path.join(MATERIALS_DIR, fname)
        ext = os.path.splitext(fname)[1].lower()
        if not os.path.isfile(fpath) or ext not in READERS:
            continue
        try:
            content = READERS[ext](fpath)
            materials.append(f"=== {fname} ===\n{content}")
            print(f"[AI] 已加载材料: {fname} ({len(content)} 字)")
        except Exception as e:
            print(f"[AI] 加载 {fname} 失败: {e}")

    if not materials:
        print("[AI] 警告: materials/ 目录为空，请放入论文/资料文件 (txt/pdf/pptx/docx)")

    return "\n\n".join(materials)


def is_question(text: str) -> bool:
    """判断一段文字是否是问题"""
    # 包含问号
    for p in QUESTION_PUNCTUATION:
        if p in text:
            return True
    # 包含疑问关键词
    for kw in QUESTION_KEYWORDS:
        if kw in text:
            return True
    return False


def start_ai_responder(text_queue: queue.Queue, answer_queue: queue.Queue) -> threading.Thread:
    """启动 AI 回答线程，从 text_queue 读取文字，将回答放入 answer_queue"""
    client = Anthropic(api_key=CLAUDE_API_KEY)
    materials = load_materials()

    system_prompt = f"""你是一个学术助手，正在帮助学生准备组会/学术讨论。

学生正在参加腾讯会议组会，老师可能提出各种问题。你需要根据学生提供的论文/资料，给出简洁、准确、有条理的回答建议。

要求：
1. 回答要简洁，适合口头表达（不要写太长，3-5句话为宜）
2. 如果资料中有直接相关内容，引用具体数据/结论
3. 如果资料中没有直接相关内容，基于你的知识给出合理回答，并标注"（资料中未找到相关内容，以下为通用回答）"
4. 用中文回答
5. 如果识别到的文字不是问题（只是陈述句），回复"不是问题，跳过"

以下是学生的论文/资料：
{materials}"""

    def _responder_loop():
        while True:
            try:
                text = text_queue.get(timeout=1.0)
            except queue.Empty:
                continue

            # 先判断是否是问题
            if not is_question(text):
                continue

            print(f"[AI] 检测到问题: {text}")

            try:
                response = client.messages.create(
                    model=CLAUDE_MODEL,
                    max_tokens=512,
                    system=system_prompt,
                    messages=[{"role": "user", "content": f"老师的问题：{text}"}],
                )
                # 兼容 MiMo 等返回 thinking+text 的格式
                answer = ""
                for block in response.content:
                    if hasattr(block, "text") and block.text:
                        answer = block.text
                        break
                if not answer:
                    answer = "[未获取到回答]"
                answer_queue.put((text, answer))
                print(f"[AI] 回答: {answer[:80]}...")
            except Exception as e:
                answer_queue.put((text, f"[API错误] {e}"))
                print(f"[AI] 错误: {e}")

    thread = threading.Thread(target=_responder_loop, daemon=True, name="AIResponder")
    thread.start()
    print("[AI] 回答线程已启动")
    return thread
